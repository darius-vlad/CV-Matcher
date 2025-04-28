from docx import Document
import glob

from scripts.service.cv_service import add_cvs
from scripts.service.job_service import add_jobs

def extract_text_from_docx(file_path):
    """Extract text from a .docx file, including paragraphs and tables."""
    doc = Document(file_path)
    text = []

    # Extract paragraphs
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text.append(paragraph.text)

    return '\n'.join(text)
def read_folder(folder_path, elements_num = -1):
    folder_path = f'{folder_path}/*.docx'

    # Extract text from all .docx files in the folder
    text_array = []
    ids_array = []
    cur_id = 1
    for file_path in glob.glob(folder_path):
        text = extract_text_from_docx(file_path)
        text_array.append(text)
        ids_array.append(cur_id)
        cur_id += 1
        if cur_id != -1 and cur_id > elements_num:
            break

    return ids_array, text_array
def test_save_jobs():
    _, job_texts = read_folder('../DataSet/job', 2)
    add_jobs(job_texts)
def test_save_cvs():
    _, cv_texts = read_folder('../DataSet/cv', 3)
    add_cvs(cv_texts)

# test_save_cvs()
test_save_jobs()
